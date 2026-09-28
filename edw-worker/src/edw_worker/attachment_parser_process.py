"""Untrusted bytes enter only this networkless, credentialless subprocess.

stdin: JSON with kind and base64 bytes. stdout: bounded page report; no paths or URLs.
"""
import base64
import ctypes
import ctypes.util
import errno
from hashlib import sha256
from io import BytesIO
import json
import sys
import zipfile
from xml.etree import ElementTree

MAX_UNPACKED = 40 * 1024 * 1024


class UnsafeFile(ValueError):
    pass


def safe_zip(raw, kind):
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        names = archive.namelist()
        if len(names) > 500 or len(set(names)) != len(names) or sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED:
            raise ValueError("unsafe archive")
        for info in archive.infolist():
            name = info.filename
            if (name.startswith("/") or "\\" in name or "\x00" in name or
                any(part in ("..", "") for part in name.split("/" ) if part != name.split("/")[-1]) or
                any(part == ".." for part in name.split("/")) or
                (info.external_attr >> 16) & 0o170000 == 0o120000 or
                info.flag_bits & 1 or info.compress_size == 0 and info.file_size > 0 or
                info.compress_size and info.file_size > info.compress_size * 100):
                raise ValueError("unsafe archive")
            if name.lower().endswith(("vbaproject.bin", ".bin", ".vbs", ".exe")):
                raise ValueError("active payload")
            if kind == 'xlsx' and name.startswith('xl/worksheets/') and name.endswith('.xml') and b'<f' in archive.read(info):
                raise ValueError('formula workbook')
            if name.endswith(".rels"):
                root = ElementTree.fromstring(archive.read(info))
                if any(node.attrib.get("TargetMode") == "External" or
                       node.attrib.get("Target", "").startswith(("http:", "https:", "file:"))
                       for node in root):
                    raise ValueError("external relationship")
        if kind == "docx" and "word/document.xml" not in names or kind == "xlsx" and "xl/workbook.xml" not in names:
            raise ValueError("wrong office format")


def page(number, text, *, section="", ocr=False):
    text = " ".join(text.split())
    if not text:
        return failed(number, "OCR_UNAVAILABLE" if ocr else "PAGE_FAILED")
    if len(text) > 200 * 4000:
        return failed(number, "RESOURCE_LIMIT")
    chunks = []
    for start in range(0, len(text), 4000):
        part = text[start:start + 4000]
        chunks.append({"text": part, "locator": f"page:{number}:char:{start}",
                       "sha256": sha256(part.encode()).hexdigest()})
    return {"number": number, "section": section[:256], "state": "OCR" if ocr else "PARSED",
            "error": None, "chunks": chunks}


def failed(number, error="PAGE_FAILED"):
    return {"number": number, "section": "", "state": "FAILED", "error": error, "chunks": []}


def ocr_image(image):
    # No model downloads: weights must already be baked into the isolated image.
    from pathlib import Path
    if not all(all(Path('/opt/eid-ocr', part, name).is_file() for name in ('inference.pdmodel', 'inference.pdiparams'))
               for part in ('det', 'rec')):
        raise ImportError('OCR model unavailable')
    from paddleocr import PaddleOCR
    engine = PaddleOCR(use_angle_cls=False, lang="ch", show_log=False,
                       det_model_dir="/opt/eid-ocr/det", rec_model_dir="/opt/eid-ocr/rec",
                       cls_model_dir="/opt/eid-ocr/cls", use_gpu=False)
    from PIL import Image
    import numpy as np
    with Image.open(BytesIO(image)) as source:
        if source.width * source.height > 16_000_000:
            raise ValueError("oversized image")
        result = engine.ocr(np.asarray(source.convert("RGB")), cls=False)
    return " ".join(str(item[1][0]) for block in result if block for item in block)


def parse(kind, raw):
    pages = []
    if kind in ("docx", "xlsx"):
        safe_zip(raw, kind)
    if kind == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(raw), strict=True)
        if any(name in reader.trailer['/Root'] for name in ('/OpenAction','/AA','/AcroForm','/Names')) or b'/EmbeddedFile' in raw or b'/JavaScript' in raw:
            raise ValueError('active PDF')
        if reader.is_encrypted or len(reader.pages) > 500:
            raise ValueError("encrypted or oversized PDF")
        for i in range(len(reader.pages)):
            try:
                part = reader.pages[i]
                # Links, annotations and embedded file actions are not published.
                if any(name in part for name in ('/AA','/OpenAction','/Annots','/RichMedia')):
                    raise UnsafeFile("active PDF")
                text = part.extract_text() or ""
                if text.strip():
                    pages.append(page(i + 1, text))
                else:
                    try:
                        import pypdfium2 as pdfium
                        rendered = pdfium.PdfDocument(raw)[i].render(scale=1.5).to_pil()
                        output = BytesIO()
                        rendered.save(output, format='PNG')
                        pages.append(page(i + 1, ocr_image(output.getvalue()), ocr=True))
                    except (ImportError, OSError):
                        pages.append(failed(i + 1, 'OCR_UNAVAILABLE'))
                    except Exception:
                        pages.append(failed(i + 1))
            except UnsafeFile:
                raise
            except Exception:
                pages.append(failed(i + 1))
    elif kind == "docx":
        from docx import Document
        document = Document(BytesIO(raw))
        lines = [paragraph.text for paragraph in document.paragraphs]
        lines += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        for i in range(0, max(1, len(lines)), 100):
            pages.append(page(i // 100 + 1, "\n".join(lines[i:i + 100]), section=f"paragraph:{i + 1}"))
    elif kind == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                if len(pages) >= 500:
                    raise ValueError("too many sheets")
                try:
                    rows = []
                    for row in sheet.iter_rows(max_row=2000, max_col=100):
                        if any(cell.data_type == "f" for cell in row):
                            raise ValueError("formula")
                        rows.append(" | ".join(str(cell.value or "")[:256] for cell in row))
                    pages.append(page(len(pages) + 1, "\n".join(rows), section=sheet.title))
                except Exception:
                    pages.append(failed(len(pages) + 1))
        finally:
            workbook.close()
    elif kind in ("png", "jpeg"):
        try:
            pages.append(page(1, ocr_image(raw), ocr=True))
        except (ImportError, OSError):
            pages.append(failed(1, "OCR_UNAVAILABLE"))
        except Exception:
            pages.append(failed(1))
    else:
        raise ValueError("unsupported type")
    return pages


def isolate_syscalls():
    """Fail closed unless a seccomp deny-list can be loaded inside the bwrap namespace."""
    path = ctypes.util.find_library("seccomp")
    if not path:
        raise RuntimeError("seccomp unavailable")
    lib = ctypes.CDLL(path)
    lib.seccomp_init.restype = ctypes.c_void_p
    lib.seccomp_init.argtypes = [ctypes.c_uint32]
    lib.seccomp_syscall_resolve_name.restype = ctypes.c_int
    lib.seccomp_syscall_resolve_name.argtypes = [ctypes.c_char_p]
    lib.seccomp_rule_add.argtypes = [ctypes.c_void_p, ctypes.c_uint32, ctypes.c_int, ctypes.c_uint]
    lib.seccomp_load.argtypes = [ctypes.c_void_p]
    lib.seccomp_release.argtypes = [ctypes.c_void_p]
    context = lib.seccomp_init(0x7fff0000)  # SCMP_ACT_ALLOW
    if not context:
        raise RuntimeError("seccomp unavailable")
    try:
        for name in (b"socket", b"socketpair", b"connect", b"bind", b"listen", b"accept", b"accept4",
                     b"execve", b"execveat", b"ptrace", b"mount", b"umount2", b"setns", b"unshare"):
            number = lib.seccomp_syscall_resolve_name(name)
            if number < 0 or lib.seccomp_rule_add(context, 0x00050000 | errno.EPERM, number, 0):
                raise RuntimeError("seccomp rule unavailable")
        if lib.seccomp_load(context):
            raise RuntimeError("seccomp load failed")
    finally:
        lib.seccomp_release(context)


def main():
    try:
        request = json.loads(sys.stdin.buffer.read(14 * 1024 * 1024))
        raw = base64.b64decode(request["bytes"], validate=True)
        if not 0 < len(raw) <= 10 * 1024 * 1024:
            raise ValueError("size")
        isolate_syscalls()
        result = {"pages": parse(request["kind"], raw)}
    except (ValueError, KeyError, zipfile.BadZipFile, ElementTree.ParseError):
        result = {"error": "UNSAFE_CONTENT"}
    except Exception:
        result = {"error": "PARSER_FAILED"}
    sys.stdout.write(json.dumps(result, ensure_ascii=True, separators=(",", ":")))


if __name__ == "__main__":
    main()
